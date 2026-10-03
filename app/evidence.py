"""Evidence tracking, audit-readiness checks and Evidence List / Doc Mapping export."""
import re
from datetime import date

from openpyxl import Workbook

from . import db, settings
from .catalog import DOCUMENTS, EVIDENCE_STATUS_TH, STAGE_BY_KEY
from .generator import items_for
from .render_xlsx import write_sheet


def overview(pid: int) -> list[dict]:
    """One row per catalog document with its evidence status and generated files."""
    project = db.get_project(pid)
    ev = {e["doc_key"]: e for e in db.rows("SELECT * FROM evidence WHERE project_id=?", (pid,))}
    docs = db.rows("SELECT d.*, u.name AS created_by_name FROM documents d LEFT JOIN users u ON u.id=d.created_by "
                   "WHERE d.project_id=? ORDER BY d.id DESC", (pid,))
    out = []
    for d in DOCUMENTS:
        e = ev.get(d["key"], {})
        files = [x for x in docs if x["doc_key"] == d["key"]]
        out.append({
            **{k: d[k] for k in ("key", "seq", "name", "name_th", "kind", "stage", "owner", "deliver", "code")},
            "val": bool(d.get("val")),
            "per_item": d.get("per_item"),
            "items": items_for(d, project["facts"]),
            "stage_title": STAGE_BY_KEY[d["stage"]]["title"],
            "status": e.get("status", "missing"),
            "links": e.get("links", ""),
            "remark": e.get("remark", ""),
            "files": files,
        })
    return out


def update(pid: int, doc_key: str, status=None, links=None, remark=None):
    db.execute("INSERT OR IGNORE INTO evidence(project_id, doc_key) VALUES (?,?)", (pid, doc_key))
    sets, args = [], []
    for col, val in (("status", status), ("links", links), ("remark", remark)):
        if val is not None:
            sets.append(f"{col}=?")
            args.append(val)
    if sets:
        db.execute(f"UPDATE evidence SET {', '.join(sets)}, updated_at=datetime('now','localtime') "
                   "WHERE project_id=? AND doc_key=?", (*args, pid, doc_key))


def checks(pid: int) -> list[dict]:
    """Audit-readiness findings: [{level, doc, message}]."""
    project = db.get_project(pid)
    cfg = settings.load_config()
    own = f"{cfg['company_short']}_{project['code']}"
    code_re = re.compile(rf"{re.escape(cfg['company_short'])}_[A-Za-z0-9][\w\-]*?-\d{{4}}-\d{{3,4}}")
    findings = []
    for row in overview(pid):
        name = f"{row['seq']} {row['name']}"
        latest = row["files"][0] if row["files"] else None
        if row["status"] == "missing" and not row["links"]:
            findings.append({"level": "error", "doc": name, "message": "ยังไม่มีหลักฐาน"})
        if latest and latest["placeholders"]:
            findings.append({"level": "warn", "doc": name,
                             "message": f"ไฟล์ล่าสุด {latest['filename']} ยังมีช่อง [ต้องระบุ] {latest['placeholders']} จุด"})
        if row["val"] and row["status"] not in ("validated", "baseline"):
            findings.append({"level": "warn", "doc": name, "message": "ต้องมีหลักฐานลูกค้ายืนยัน (Val/Sign-off) จริง"})
        if row["per_item"]:
            done = {f["item_index"] for f in row["files"]}
            missing = [i + 1 for i in range(len(row["items"])) if i not in done]
            if missing:
                findings.append({"level": "warn", "doc": name,
                                 "message": f"ยังไม่ได้จัดทำเอกสารรายการที่ {', '.join(map(str, missing))}"})
        for m in code_re.findall(f"{row['links']} {row['remark']}"):
            if not m.startswith(own):
                findings.append({"level": "error", "doc": name,
                                 "message": f"พบรหัสโครงการอื่น '{m}' ในหลักฐาน — ตรวจว่าคัดลอกมาจากโครงการเก่าหรือไม่"})
    for stage_key, gaps in project["facts"].get("_gaps", {}).items():
        for g in gaps:
            findings.append({"level": "info", "doc": STAGE_BY_KEY[stage_key]["title"], "message": g})
    return findings


def _meta(project, cfg, title_th, title_en):
    return {
        "company": cfg["company_name"], "title_th": title_th, "title_en": title_en,
        "project_name": project["name"], "project_code": project["code"], "customer": project["customer"],
        "doc_no": f"{cfg['company_short']}_{project['code']}_EvidenceList", "version": "-",
        "prepared_date": date.today().isoformat(),
    }


def export(pid: int) -> str:
    """Write Evidence List + Doc Mapping workbook; return path relative to storage."""
    project = db.get_project(pid)
    cfg = settings.load_config()
    rows = overview(pid)

    def evidence_names(r):
        # files are newest-first; keep the latest version of each item
        seen, latest = set(), []
        for f in r["files"]:
            if f["item_index"] not in seen:
                seen.add(f["item_index"])
                latest.append(f["filename"])
        return "\n".join(latest)

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Evidence List")
    write_sheet(ws, _meta(project, cfg, "รายการหลักฐาน", "Evidence List"), "",
                ["ลำดับ", "Work Products", "Evidence Name", "Source/Path in Project Repository", "สถานะ", "Remark"],
                [[r["seq"], r["name"], evidence_names(r), r["links"], EVIDENCE_STATUS_TH[r["status"]], r["remark"]]
                 for r in rows])
    ws = wb.create_sheet("Doc Mapping")
    write_sheet(ws, _meta(project, cfg, "แผนผังเอกสาร", "Doc Mapping"), "",
                ["Deliver Type", "Seq.", "ISO/IEC29110 Standard Document", "Phase", "Owner", "Link Document", "List Doc", "สถานะ"],
                [[r["deliver"], r["seq"], r["name"], r["stage_title"], r["owner"], r["links"], evidence_names(r),
                  EVIDENCE_STATUS_TH[r["status"]]] for r in rows])
    ws = wb.create_sheet("Findings")
    write_sheet(ws, _meta(project, cfg, "ผลการตรวจความพร้อม", "Audit Readiness Findings"), "",
                ["ระดับ", "เอกสาร/ขั้น", "รายละเอียด"],
                [[f["level"], f["doc"], f["message"]] for f in checks(pid)])

    out_dir = settings.STORAGE_DIR / re.sub(r"[^\w\-.]+", "_", project["code"])
    out_dir.mkdir(parents=True, exist_ok=True)
    name = f"{cfg['company_short']}_{project['code']}_EvidenceList_{date.today():%d-%b-%y}.xlsx"
    wb.save(out_dir / name)
    return (out_dir / name).relative_to(settings.STORAGE_DIR).as_posix()
