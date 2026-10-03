"""Render document JSON into .docx using Thai official-document conventions.

- TH Sarabun New 16pt (Thai + Latin), headings bold
- Margins: top 2.5 cm, bottom 2 cm, left 3 cm, right 2 cm
- Page numbers top-centre as "- n -" from page 2
- Paragraph first-line indent 2.5 cm
- Signature blocks: (ลงชื่อ) ....... / ( name ) / ตำแหน่ง / วันที่
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .thai import thai_date, thai_digits

FONT = "TH Sarabun New"
BODY_PT = 16


def configure(font: str | None = None, size: int | None = None):
    """Apply the font settings from config before rendering."""
    global FONT, BODY_PT
    FONT = font or FONT
    BODY_PT = int(size or BODY_PT)


def _set_run_font(run, size=None, bold=None):
    size = size or BODY_PT
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs")
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(size * 2)))
    if bold:
        bcs = rpr.find(qn("w:bCs"))
        if bcs is None:
            rpr.append(OxmlElement("w:bCs"))


def _base_document():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(BODY_PT * 2))
    rpr.append(szcs)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
    sec.header_distance = Cm(1.2)
    sec.different_first_page_header_footer = True
    return doc


def _page_number_header(doc, digits: bool):
    for sec in doc.sections:
        p = sec.header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("- ")
        _set_run_font(r)
        r = p.add_run()
        _set_run_font(r)
        for tag, text in (("begin", None), (None, "PAGE" if not digits else r"PAGE \* ThaiArabic"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            r._element.append(el)
        r = p.add_run(" -")
        _set_run_font(r)


def para(doc_or_cell, text="", size=None, bold=False, align=None, indent=False, space_after=0):
    p = doc_or_cell.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if indent:
        p.paragraph_format.first_line_indent = Cm(2.5)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        _set_run_font(p.add_run(str(text)), size=size, bold=bold)
    return p


def _label_value(doc, label, value, label_pt=20):
    p = doc.add_paragraph()
    _set_run_font(p.add_run(label + "  "), size=label_pt, bold=True)
    _set_run_font(p.add_run(str(value or "")))
    return p


def _shade(cell, hex_fill="D9D9D9"):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _cell_text(cell, text, bold=False, size=14, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = str(text if text is not None else "").split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        _set_run_font(r, size=size, bold=bold)
        if i < len(lines) - 1:
            r.add_break(WD_BREAK.LINE)


def table(doc, columns, rows, widths_cm=None, header_fill="D9D9D9"):
    columns = [str(c) for c in (columns or [])]
    if not columns:
        return None
    t = doc.add_table(rows=1, cols=len(columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(columns):
        _cell_text(t.rows[0].cells[i], c, bold=True, align="center")
        _shade(t.rows[0].cells[i], header_fill)
    if not rows:  # empty form table: leave lines to write on
        rows = [[""] * len(columns) for _ in range(3)]
    for r in rows:
        cells = t.add_row().cells
        r = list(r) if isinstance(r, (list, tuple)) else [r]
        for i in range(len(columns)):
            _cell_text(cells[i], r[i] if i < len(r) else "")
    if not widths_cm:
        widths_cm = _auto_widths(columns, rows)
    t.autofit = False
    for row in t.rows:
        for i, w in enumerate(widths_cm[: len(columns)]):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def _auto_widths(columns, rows, total_cm=16.0, min_cm=2.0):
    """Share the text width by content length so long-text columns get room."""
    weights = []
    for i, c in enumerate(columns):
        cells = [str(c)] + [str(r[i]) for r in (rows or []) if isinstance(r, (list, tuple)) and i < len(r)]
        longest_line = max((len(line) for cell in cells for line in cell.split("\n")), default=4)
        avg = sum(len(x) for x in cells) / len(cells)
        header_len = max(len(w) for w in str(c).split()) if str(c).strip() else 4
        weights.append(max(4.0, header_len * 1.6, min(60.0, (longest_line + avg) / 2)))
    total_w = sum(weights)
    widths = [max(min_cm, total_cm * w / total_w) for w in weights]
    scale = total_cm / sum(widths)
    return [w * scale for w in widths]


def _normalize_block(b) -> dict:
    """The model doesn't always follow the block schema; coerce what it sends."""
    if isinstance(b, dict):
        b = dict(b)
        # synonyms the model uses instead of the schema's field names
        for alias, field in (("content", "text"), ("headers", "columns"), ("title", "caption"), ("data", "rows")):
            if alias in b and field not in b:
                b[field] = b.pop(alias)
        if isinstance(b.get("text"), list):
            b = {"type": "bullets", "items": b["text"]}
        rows = b.get("rows")
        if isinstance(rows, list) and rows and all(isinstance(x, dict) for x in rows):
            cols = b.get("columns") or list(dict.fromkeys(k for x in rows for k in x))
            b.update(type="table", columns=cols, rows=[[x.get(c, "") for c in cols] for x in rows])
        typ = b.get("type")
        if typ in ("paragraph", "bullets", "numbered", "table", "heading"):
            return b
        if "columns" in b or "rows" in b:
            return {**b, "type": "table"}
        data = b.get("rows") or b.get("items")
        if isinstance(data, list) and data and all(isinstance(x, dict) for x in data):
            cols = list(dict.fromkeys(k for x in data for k in x))
            return {"type": "table", "caption": b.get("caption") or b.get("title", ""),
                    "columns": cols, "rows": [[x.get(c, "") for c in cols] for x in data]}
        if isinstance(data, list):
            return {**b, "type": "numbered" if typ in ("ordered", "numbered_list") else "bullets", "items": data}
        return {"type": "paragraph", "text": b.get("text") or "; ".join(f"{k}: {v}" for k, v in b.items())}
    if isinstance(b, (list, tuple)):
        return {"type": "bullets", "items": [x if isinstance(x, str) else str(x) for x in b]}
    return {"type": "paragraph", "text": "" if b is None else str(b)}


def _sections(content: dict) -> list[dict]:
    out = []
    for s in content.get("sections") or []:
        if isinstance(s, dict):
            out.append(s)
        else:
            out.append({"heading": "", "blocks": s if isinstance(s, list) else [s]})
    return out


def _blocks(doc, blocks):
    if isinstance(blocks, (str, dict)):
        blocks = [blocks]
    for b in blocks or []:
        b = _normalize_block(b)
        typ = b.get("type", "paragraph")
        if typ in ("bullets", "numbered"):
            for i, item in enumerate(b.get("items") or [], 1):
                p = para(doc, ("%d) " % i if typ == "numbered" else "• ") + str(item))
                p.paragraph_format.left_indent = Cm(1.5)
                p.paragraph_format.first_line_indent = Cm(-0.6)
        elif typ == "heading":
            para(doc, b.get("text", ""), bold=True, space_after=2)
        elif typ == "table":
            if b.get("caption"):
                para(doc, b["caption"], bold=True)
            table(doc, b.get("columns"), b.get("rows"))
        else:
            # one paragraph per line so every line gets the same indent (e.g. ☐ choices)
            for line in str(b.get("text", "")).split("\n"):
                para(doc, line, indent=True, align="justify")


def _end_paragraph(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    _set_run_font(p.add_run(""), size=1)
    rpr = p._p.get_or_add_pPr()
    mark = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    mark.append(sz)
    rpr.append(mark)


def _signatures(doc, sigs, digits=False):
    sigs = [s for s in (sigs or []) if isinstance(s, dict)]
    if not sigs:
        return
    doc.add_paragraph()
    cols = 2 if len(sigs) > 1 else 1
    nrows = (len(sigs) + cols - 1) // cols
    t = doc.add_table(rows=nrows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t.rows:  # never split a signature row across pages
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
    for idx, s in enumerate(sigs):
        cell = t.rows[idx // cols].cells[idx % cols]
        cell.text = ""
        lines = [
            s.get("label", ""),
            "(ลงชื่อ)......................................................",
            f"( {s.get('name') or '......................................................'} )",
            f"ตำแหน่ง {s.get('position') or '..........................................'}",
            "วันที่ ........../........................./..............",
        ]
        first = True
        for ln in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            if first and idx >= cols:
                p.paragraph_format.space_before = Pt(24)
            if ln.startswith("(ลงชื่อ)"):
                p.paragraph_format.space_before = Pt(14)
            first = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_run_font(p.add_run(ln), bold=(ln == s.get("label")))


def _cover(doc, meta, digits):
    for _ in range(4):
        doc.add_paragraph()
    para(doc, meta["company"], size=20, bold=True, align="center")
    doc.add_paragraph()
    para(doc, meta["title_th"], size=28, bold=True, align="center")
    para(doc, meta["title_en"], size=20, align="center")
    doc.add_paragraph()
    para(doc, f"โครงการ {meta['project_name']}", size=20, bold=True, align="center")
    para(doc, f"ลูกค้า {meta['customer']}", size=18, align="center")
    for _ in range(6):
        doc.add_paragraph()
    v = meta["version"]
    para(doc, f"รหัสเอกสาร {meta['doc_no']}", align="center")
    para(doc, f"ฉบับที่ {thai_digits(v) if digits else v}", align="center")
    para(doc, f"วันที่จัดทำเอกสาร {thai_date(meta['prepared_date'], digits)}", align="center")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _doc_control(doc, meta, content):
    para(doc, "ข้อมูลเอกสาร", size=18, bold=True)
    table(doc, ["รายการ", "รายละเอียด"], [
        ["ชื่อเอกสาร", f"{meta['title_th']} ({meta['title_en']})"],
        ["รหัสเอกสาร", meta["doc_no"]],
        ["โครงการ", f"{meta['project_name']} ({meta['project_code']})"],
        ["อ้างอิงมาตรฐาน", "ISO/IEC 29110 Basic Profile"],
        ["วันที่จัดทำเอกสาร", thai_date(meta["prepared_date"])],
        ["วันที่ของเหตุการณ์/ข้อมูล", content.get("event_date_text") or "-"],
    ], widths_cm=[5, 11])
    para(doc, "ประวัติการแก้ไขเอกสาร", size=18, bold=True)
    table(doc, ["ฉบับที่", "วันที่", "รายละเอียดการแก้ไข", "ผู้จัดทำ"],
          [[meta["version"], thai_date(meta["prepared_date"]), "จัดทำเอกสารฉบับแรก", meta.get("prepared_by") or ""]],
          widths_cm=[2, 4, 7, 3])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def render_report(path, meta, content, digits=False):
    doc = _base_document()
    if meta.get("cover", True):
        _cover(doc, meta, digits)
    _doc_control(doc, meta, content)
    for i, s in enumerate(_sections(content), 1):
        heading = s.get("heading", "")
        para(doc, heading if heading[:1].isdigit() else f"{i}. {heading}", size=18, bold=True, space_after=4)
        _blocks(doc, s.get("blocks"))
        doc.add_paragraph()
    _signatures(doc, content.get("signatures"), digits)
    _end_paragraph(doc)
    _page_number_header(doc, digits)
    doc.save(path)


def render_form(path, meta, content, digits=False):
    doc = _base_document()
    para(doc, meta["company"], size=16, bold=True, align="center")
    para(doc, meta["title_th"], size=24, bold=True, align="center")
    para(doc, meta["title_en"], size=16, align="center", space_after=6)
    fields = content.get("fields") or []
    base = [["โครงการ", f"{meta['project_name']} ({meta['project_code']})"],
            ["ลูกค้า", meta["customer"]],
            ["รหัสเอกสาร", meta["doc_no"]],
            ["วันที่จัดทำเอกสาร", thai_date(meta["prepared_date"], digits)]]
    base += [[f.get("label", ""), f.get("value", "")] for f in fields if isinstance(f, dict)]
    table(doc, ["รายการ", "รายละเอียด"], base, widths_cm=[5, 11])
    for i, s in enumerate(_sections(content), 1):
        para(doc, f"{i}. {s.get('heading', '')}", size=18, bold=True, space_after=4)
        _blocks(doc, s.get("blocks"))
    _signatures(doc, content.get("signatures"), digits)
    _end_paragraph(doc)
    _page_number_header(doc, digits)
    doc.save(path)


def render_memo(path, meta, content, digits=False):
    """บันทึกข้อความ (internal memo) layout."""
    doc = _base_document()
    m = content.get("memo") or {}
    para(doc, "บันทึกข้อความ", size=29, bold=True, align="center", space_after=6)
    _label_value(doc, "ส่วนงาน", f"{meta['department']} {meta['company']}")
    p = doc.add_paragraph()
    _set_run_font(p.add_run("ที่  "), size=20, bold=True)
    _set_run_font(p.add_run(f"{m.get('ref_no') or meta['doc_no']}" + " " * 20))
    _set_run_font(p.add_run("วันที่  "), size=20, bold=True)
    _set_run_font(p.add_run(m.get("date_text") or "......../................/........"))
    _label_value(doc, "เรื่อง", m.get("subject") or meta["title_th"])
    # horizontal rule
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)
    para(doc, f"เรียน  {m.get('to') or '..............................'}", space_after=6)
    body = m.get("body") or []
    for text in ([body] if isinstance(body, str) else body):
        para(doc, text, indent=True, align="justify", space_after=6)
    for s in _sections(content):
        if s.get("heading"):
            para(doc, s["heading"], bold=True)
        _blocks(doc, s.get("blocks"))
    _signatures(doc, content.get("signatures") or [{"label": "", "name": "", "position": ""}], digits)
    _end_paragraph(doc)
    _page_number_header(doc, digits)
    doc.save(path)
