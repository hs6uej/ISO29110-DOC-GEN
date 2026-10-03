"""Summarise the intent middleware logs into a test-evidence workbook.

Reads the hourly logs written by the middleware (logs/YYYY/MM/DD/HH/logYYYYMMDDHH.txt)
and produces an .xlsx with: summary figures, one row per request, errors, intents.
Every figure comes from the log lines; nothing is estimated.

usage: python tools/log_summary.py <logs_dir> <out.xlsx> [--from YYYY-MM-DD] [--to YYYY-MM-DD] [--slow 3.0]
"""
import argparse
import re
import statistics
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

LINE_RE = re.compile(r"^(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}),\d+ - (\w+) - (.*)$")
AUDIO_RE = re.compile(r"Audio URL:\s*(\S+)")
GRAM_RE = re.compile(r"Gram URL:\s*(\S+)")
GEMINI_RE = re.compile(r"Gemini .*?([\d.]+) วิ")
RESULT_RE = re.compile(r"Intent='(.*?)', Confidence=([\d.]+), Text='(.*)'")
END_OK_RE = re.compile(r"\[END\] ส่ง Response กลับสำเร็จ \(เวลารวมทั้งหมด ([\d.]+) วิ\)")
END_ERR_RE = re.compile(r"\[END\] เกิด Error ส่ง NULL กลับ \(เวลารวม ([\d.]+) วิ\)")
ERROR_RE = re.compile(r"ระบบขัดข้อง:\s*(.*)")
UTT_RE = re.compile(r"-(utt\d+)-")


def parse(log_dir: Path, date_from=None, date_to=None):
    requests, cur = [], None
    for path in sorted(log_dir.rglob("log*.txt")):
        for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
            m = LINE_RE.match(raw)
            if not m:
                continue
            ts, level, msg = m.groups()
            day = ts[:10]
            if (date_from and day < date_from) or (date_to and day > date_to):
                continue
            if "[START]" in msg:
                cur = {"start": ts, "audio": "", "gram": "", "intent": "", "confidence": None, "text": "",
                       "gemini_s": None, "total_s": None, "status": "", "error": "", "file": path.name}
                requests.append(cur)
                continue
            if cur is None:
                continue
            if (a := AUDIO_RE.search(msg)):
                cur["audio"] = a.group(1)
            elif (g := GRAM_RE.search(msg)):
                cur["gram"] = g.group(1)
            elif "Gemini" in msg and (gm := GEMINI_RE.search(msg)):
                cur["gemini_s"] = float(gm.group(1))
            elif (r := RESULT_RE.search(msg)):
                cur["intent"], cur["confidence"], cur["text"] = r.group(1), float(r.group(2)), r.group(3)
            elif level == "ERROR" and (e := ERROR_RE.search(msg)):
                cur["error"] = e.group(1)
            elif (ok := END_OK_RE.search(msg)):
                cur["total_s"], cur["status"] = float(ok.group(1)), "OK"
                cur = None
            elif (err := END_ERR_RE.search(msg)):
                cur["total_s"], cur["status"] = float(err.group(1)), "NULL (fallback)"
                cur = None
    return requests


def p95(values):
    if not values:
        return None
    values = sorted(values)
    return values[min(len(values) - 1, int(round(0.95 * (len(values) - 1))))]


def error_kind(msg: str) -> str:
    for key, label in (("unknown intent", "intent ไม่อยู่ใน prompt"), ("not allowed", "host นอก whitelist"),
                       ("404", "ไม่พบไฟล์ (404)"), ("Timeout", "timeout"), ("timed out", "timeout"),
                       ("exceeds max size", "ไฟล์ใหญ่เกิน"), ("empty", "ไฟล์ว่าง")):
        if key.lower() in msg.lower():
            return label
    return "อื่น ๆ"


def write(requests, out: Path, slow: float):
    done = [r for r in requests if r["status"]]
    ok = [r for r in done if r["status"] == "OK"]
    totals = [r["total_s"] for r in done if r["total_s"] is not None]
    ok_totals = [r["total_s"] for r in ok]
    days = sorted({r["start"][:10] for r in done})

    wb = Workbook()
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9D9D9")

    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("ช่วงข้อมูล", f"{days[0]} ถึง {days[-1]}" if days else "-"),
        ("จำนวน request ทั้งหมด", len(done)),
        ("สำเร็จ (คืน intent)", len(ok)),
        ("ผิดพลาด (คืน NULL → fallback Nuance)", len(done) - len(ok)),
        ("อัตราสำเร็จ (%)", round(100 * len(ok) / len(done), 2) if done else "-"),
        ("เวลารวมเฉลี่ย (วินาที)", round(statistics.mean(totals), 3) if totals else "-"),
        ("เวลารวม p95 (วินาที)", p95(totals) if totals else "-"),
        (f"สำเร็จแต่เกิน {slow} วินาที (curl ตัด → fallback)", sum(1 for t in ok_totals if t > slow)),
        (f"สำเร็จภายใน {slow} วินาที (%)", round(100 * sum(1 for t in ok_totals if t <= slow) / len(ok_totals), 2) if ok_totals else "-"),
        ("จำนวน request ที่ยังไม่จบในไฟล์ log", len(requests) - len(done)),
        ("หมายเหตุ", "ตัวเลขทั้งหมดคำนวณจากไฟล์ log ของ middleware โดยตรง"),
    ]
    for r in rows:
        ws.append(r)
    for c in ws["A"]:
        c.font = head
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 40

    ws = wb.create_sheet("Requests")
    cols = ["เวลาเริ่ม", "utterance", "intent", "confidence", "ข้อความที่ถอดได้", "เวลา Gemini (วิ)",
            "เวลารวม (วิ)", f"เกิน {slow} วิ", "สถานะ", "error", "grammar", "audio URL", "ไฟล์ log"]
    ws.append(cols)
    for r in done:
        utt = UTT_RE.search(r["audio"])
        ws.append([r["start"], utt.group(1) if utt else "", r["intent"], r["confidence"], r["text"], r["gemini_s"],
                   r["total_s"], "ใช่" if r["total_s"] and r["total_s"] > slow else "", r["status"], r["error"],
                   r["gram"].rsplit("/", 1)[-1], r["audio"], r["file"]])

    ws = wb.create_sheet("Errors")
    ws.append(["ประเภท", "จำนวน"])
    for kind, n in Counter(error_kind(r["error"]) for r in done if r["status"] != "OK").most_common():
        ws.append([kind, n])

    ws = wb.create_sheet("Intents")
    ws.append(["intent", "จำนวนครั้ง", "confidence เฉลี่ย"])
    by_intent = {}
    for r in ok:
        by_intent.setdefault(r["intent"], []).append(r["confidence"] or 0.0)
    for intent, confs in sorted(by_intent.items(), key=lambda kv: -len(kv[1])):
        ws.append([intent, len(confs), round(statistics.mean(confs), 3)])

    ws = wb.create_sheet("Daily")
    ws.append(["วันที่", "request", "สำเร็จ", "NULL", "เวลารวมเฉลี่ย (วิ)"])
    for day in days:
        d = [r for r in done if r["start"].startswith(day)]
        t = [r["total_s"] for r in d if r["total_s"] is not None]
        ws.append([day, len(d), sum(r["status"] == "OK" for r in d), sum(r["status"] != "OK" for r in d),
                   round(statistics.mean(t), 3) if t else ""])

    for sheet in wb.worksheets[1:]:
        for c in sheet[1]:
            c.font, c.fill = head, fill
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = max(10, min(70, width + 2))
    wb.save(out)
    return {"requests": len(done), "ok": len(ok), "days": len(days)}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("logs_dir", type=Path)
    ap.add_argument("out", type=Path)
    ap.add_argument("--from", dest="date_from")
    ap.add_argument("--to", dest="date_to")
    ap.add_argument("--slow", type=float, default=3.0, help="threshold in seconds (dialplan curl timeout)")
    args = ap.parse_args()
    for d in (args.date_from, args.date_to):
        if d:
            datetime.strptime(d, "%Y-%m-%d")
    stats = write(parse(args.logs_dir, args.date_from, args.date_to), args.out, args.slow)
    print(f"{stats['requests']} requests ({stats['ok']} OK) over {stats['days']} day(s) -> {args.out}")


if __name__ == "__main__":
    main()
