"""Render sheet JSON into .xlsx with a Thai official header block."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .thai import thai_date

FONT = "TH Sarabun New"


def configure(font: str | None = None):
    global FONT
    FONT = font or FONT
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")


def _safe_sheet_name(name: str, used: set) -> str:
    bad = '[]:*?/\\'
    base = "".join("_" if ch in bad else ch for ch in (name or "Sheet")).strip()[:31] or "Sheet"
    out, n = base, 2
    while out in used:
        suffix = f"_{n}"
        out = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(out)
    return out


def write_sheet(ws, meta: dict, title: str, columns: list, rows: list):
    ncols = max(len(columns or []), 2)
    header = [
        (meta["company"], 16, True),
        (f"{meta['title_th']} ({meta['title_en']})", 18, True),
        (f"โครงการ: {meta['project_name']} ({meta['project_code']})   ลูกค้า: {meta['customer']}", 14, False),
        (f"รหัสเอกสาร: {meta['doc_no']}   ฉบับที่: {meta['version']}   วันที่จัดทำ: {thai_date(meta['prepared_date'])}", 14, False),
    ]
    if title:
        header.append((title, 16, True))
    r = 1
    for text, size, bold in header:
        ws.cell(row=r, column=1, value=text).font = Font(name=FONT, size=size, bold=bold)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        r += 1
    r += 1
    head_row = r
    for c, name in enumerate(columns or [], 1):
        cell = ws.cell(row=r, column=c, value=str(name))
        cell.font = Font(name=FONT, size=14, bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows or []:
        r += 1
        row = list(row) if isinstance(row, (list, tuple)) else [row]
        for c in range(1, len(columns or []) + 1):
            v = row[c - 1] if c - 1 < len(row) else ""
            cell = ws.cell(row=r, column=c, value=v if isinstance(v, (int, float)) else str(v if v is not None else ""))
            cell.font = Font(name=FONT, size=14)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    # column widths from content length (Thai text is narrow per char)
    for c in range(1, len(columns or []) + 1):
        longest = max((len(str(ws.cell(row=i, column=c).value or "")) for i in range(head_row, r + 1)), default=10)
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(60, longest * 0.9 + 2))
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return r


def _signature_rows(ws, start_row: int, sigs: list):
    r = start_row + 2
    for s in sigs or []:
        if not isinstance(s, dict):
            continue
        for text in (s.get("label", ""),
                     "(ลงชื่อ)..............................................",
                     f"( {s.get('name') or '..............................................'} )",
                     f"ตำแหน่ง {s.get('position') or '..............................'}",
                     "วันที่ ......../................/.........."):
            ws.cell(row=r, column=2, value=text).font = Font(name=FONT, size=14, bold=(text == s.get("label")))
            r += 1
        r += 1


def render_workbook(path, meta: dict, content: dict):
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    sheets = content.get("sheets") or []
    if not sheets:
        sheets = [{"name": "Sheet1", "title": "", "columns": ["หมายเหตุ"], "rows": [["[ต้องระบุ: ไม่มีข้อมูล]"]]}]
    last_ws, last_row = None, 0
    for s in sheets:
        ws = wb.create_sheet(_safe_sheet_name(s.get("name"), used))
        last_row = write_sheet(ws, meta, s.get("title", ""), s.get("columns") or [], s.get("rows") or [])
        last_ws = ws
    if content.get("signatures") and last_ws is not None:
        _signature_rows(last_ws, last_row, content["signatures"])
    wb.save(path)
