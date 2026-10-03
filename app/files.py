"""Manage generated files: rename, delete, upload an edited copy as a new version."""
import os
import re

from docx import Document
from openpyxl import load_workbook

from . import db, settings
from .catalog import DOC_BY_KEY
from .generator import PLACEHOLDER, VERSION_RE, _clean_filename, _next_version, _slug

ALLOWED_EXT = {"docx", "xlsx"}


def get(doc_id: int) -> dict:
    d = db.row("SELECT * FROM documents WHERE id=?", (doc_id,))
    if not d:
        raise LookupError("ไม่พบไฟล์")
    return d


def _raw_json(rel_path: str):
    return settings.DATA_DIR / "content" / (os.path.splitext(os.path.basename(rel_path))[0] + ".json")


def count_placeholders(path) -> int:
    """Count [ต้องระบุ: ...] markers left in a .docx or .xlsx."""
    text = []
    if str(path).lower().endswith(".docx"):
        d = Document(path)
        text += [p.text for p in d.paragraphs]
        for t in d.tables:
            text += [c.text for row in t.rows for c in row.cells]
    else:
        wb = load_workbook(path, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                text += [str(v) for v in row if v is not None]
    return "\n".join(text).count(PLACEHOLDER)


def _sync_evidence(project_id: int, doc_key: str):
    """No files left for a draft slot → back to 'missing'."""
    left = db.row("SELECT COUNT(*) n FROM documents WHERE project_id=? AND doc_key=?", (project_id, doc_key))["n"]
    if not left:
        db.execute("UPDATE evidence SET status='missing' WHERE project_id=? AND doc_key=? AND status='draft'",
                   (project_id, doc_key))


def rename(doc_id: int, filename: str | None = None, version: str | None = None) -> dict:
    d = get(doc_id)
    old_path = settings.STORAGE_DIR / d["rel_path"]
    ext = old_path.suffix.lower()
    new_name = d["filename"]
    if filename:
        new_name = _clean_filename(filename)
        if not new_name.lower().endswith(ext):
            new_name = re.sub(r"\.(docx|xlsx)$", "", new_name, flags=re.I) + ext
    if version is not None and not VERSION_RE.match(version.strip()):
        raise ValueError("รูปแบบฉบับที่ต้องเป็นตัวเลข เช่น 1.0 หรือ 2.1")
    new_path = old_path.with_name(new_name)
    if new_path != old_path:
        if new_path.exists():
            raise ValueError(f"มีไฟล์ชื่อ {new_name} อยู่แล้ว")
        if old_path.exists():
            os.replace(old_path, new_path)
        raw = _raw_json(d["rel_path"])
        if raw.exists():
            os.replace(raw, raw.with_name(new_path.stem + ".json"))
    rel = new_path.relative_to(settings.STORAGE_DIR).as_posix()
    db.execute("UPDATE documents SET filename=?, rel_path=?, version=? WHERE id=?",
               (new_name, rel, (version or d["version"]).strip(), doc_id))
    return get(doc_id)


def delete(doc_id: int) -> dict:
    d = get(doc_id)
    for p in (settings.STORAGE_DIR / d["rel_path"], _raw_json(d["rel_path"])):
        p.unlink(missing_ok=True)
    db.execute("DELETE FROM documents WHERE id=?", (doc_id,))
    _sync_evidence(d["project_id"], d["doc_key"])
    return {"ok": True, "filename": d["filename"]}


def upload_version(base_id: int, original_name: str, content: bytes, version: str | None = None) -> dict:
    """Store an edited copy of an existing file as the next version of the same document/item."""
    base = get(base_id)
    ext = os.path.splitext(original_name or "")[1].lower().lstrip(".")
    base_ext = os.path.splitext(base["filename"])[1].lower().lstrip(".")
    if ext not in ALLOWED_EXT:
        raise ValueError("อัปโหลดได้เฉพาะไฟล์ .docx หรือ .xlsx")
    if ext != base_ext:
        raise ValueError(f"ไฟล์ต้องเป็น .{base_ext} เหมือนต้นฉบับ")
    cfg = settings.load_config()
    version = (version or _next_version(base["project_id"], base["doc_key"], base["item_index"], cfg)).strip()
    if not VERSION_RE.match(version):
        raise ValueError("รูปแบบฉบับที่ต้องเป็นตัวเลข เช่น 1.0 หรือ 2.1")
    # same name as the base file with the version swapped; fall back to appending it
    stem = os.path.splitext(base["filename"])[0]
    stem = re.sub(rf"v{re.escape(base['version'])}$", f"v{version}", stem) if stem.endswith(f"v{base['version']}") \
        else f"{stem}_v{version}"
    new_name = _clean_filename(stem) + "." + ext
    project = db.get_project(base["project_id"])
    stage = DOC_BY_KEY[base["doc_key"]]["stage"] if base["doc_key"] in DOC_BY_KEY else "other"
    out_dir = settings.STORAGE_DIR / _slug(project["code"]) / stage
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / new_name
    if path.exists():
        raise ValueError(f"มีไฟล์ชื่อ {new_name} อยู่แล้ว ระบุฉบับที่อื่น")
    path.write_bytes(content)
    try:
        placeholders = count_placeholders(path)
    except Exception:
        path.unlink(missing_ok=True)
        raise ValueError("เปิดไฟล์ไม่ได้ ไฟล์อาจเสียหรือไม่ใช่ Word/Excel")
    rel = path.relative_to(settings.STORAGE_DIR).as_posix()
    doc_id = db.execute(
        "INSERT INTO documents(project_id, doc_key, item_index, filename, rel_path, version, placeholders) VALUES (?,?,?,?,?,?,?)",
        (base["project_id"], base["doc_key"], base["item_index"], new_name, rel, version, placeholders),
    )
    return get(doc_id)
